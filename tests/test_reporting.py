from __future__ import annotations

from pathlib import Path

import pytest

from hexharness.evidence.findings import Severity
from hexharness.evidence.store import EvidenceStore
from hexharness.providers.fake import FakeProvider
from hexharness.providers.types import ModelResponse, StopReason, TextBlock
from hexharness.reporting.generate import fill_generative
from hexharness.reporting.model import ReportModel
from hexharness.reporting.render import introspect_placeholders, render, validate_template

DEFAULT_TEMPLATE = Path(__file__).resolve().parents[1] / "hexharness" / "reporting" / "templates" / "default.html.j2"

META = {"name": "Acme Web App", "client": "Acme", "date": "2026-10-06", "scope_summary": "acme.example"}


async def _seeded_store() -> EvidenceStore:
    store = EvidenceStore(":memory:")
    crit = await store.add_candidate(title="SQLi in login", severity=Severity.CRITICAL, target="acme.example")
    high = await store.add_candidate(title="Reflected XSS", severity=Severity.HIGH, target="acme.example/search")
    med = await store.add_candidate(title="Missing HSTS", severity=Severity.MEDIUM)
    noise = await store.add_candidate(title="Not real", severity=Severity.LOW)
    await store.confirm(crit.id, curator="tester")
    await store.confirm(high.id, curator="tester")
    await store.confirm(med.id, curator="tester")
    await store.reject(noise.id, curator="tester", reason="false positive")
    # candidate left untouched -> must NOT appear in the report
    await store.add_candidate(title="Unreviewed", severity=Severity.HIGH)
    return store


async def test_from_evidence_uses_only_confirmed():
    store = await _seeded_store()
    model = ReportModel.from_evidence(META, store)

    titles = {f.title for f in model.findings}
    assert titles == {"SQLi in login", "Reflected XSS", "Missing HSTS"}
    assert "Not real" not in titles  # rejected
    assert "Unreviewed" not in titles  # candidate

    assert model.severity_counts == {
        "critical": 1, "high": 1, "medium": 1, "low": 0, "info": 0,
    }
    # highest severity first
    assert [f.severity for f in model.findings][:3] == ["critical", "high", "medium"]
    assert model.engagement_name == "Acme Web App"


async def test_validate_template_catches_unknown_field(tmp_path):
    pytest.importorskip("jinja2")  # HTML introspection needs Jinja
    bad = tmp_path / "bad.html.j2"
    bad.write_text("<p>{{ engagement_name }} {{ bogus_field }}</p>")
    with pytest.raises(ValueError, match="bogus_field"):
        validate_template(bad)


def test_default_template_validates():
    pytest.importorskip("jinja2")
    placeholders = introspect_placeholders(DEFAULT_TEMPLATE)
    # every placeholder is a real ReportModel field
    assert placeholders <= set(ReportModel.model_fields)
    validate_template(DEFAULT_TEMPLATE)  # must not raise


async def test_render_html(tmp_path):
    pytest.importorskip("jinja2")
    store = await _seeded_store()
    model = ReportModel.from_evidence(META, store)
    out = tmp_path / "report.html"
    render(model, DEFAULT_TEMPLATE, out)
    html = out.read_text()
    assert "SQLi in login" in html
    assert "Acme Web App" in html
    assert "Not real" not in html


async def test_fill_generative_with_fake_provider():
    store = await _seeded_store()
    model = ReportModel.from_evidence(META, store)
    provider = FakeProvider([
        ModelResponse(content=[TextBlock(text="Three issues found.")], stop_reason=StopReason.END_TURN),
        ModelResponse(content=[TextBlock(text="Material business risk.")], stop_reason=StopReason.END_TURN),
    ])
    await fill_generative(model, provider)
    assert model.executive_summary == "Three issues found."
    assert model.risk_narrative == "Material business risk."
    assert len(provider.calls) == 2


async def test_render_pdf_weasyprint(tmp_path):
    pytest.importorskip("jinja2")
    pytest.importorskip("weasyprint")
    store = await _seeded_store()
    model = ReportModel.from_evidence(META, store)
    out = tmp_path / "report.pdf"
    render(model, DEFAULT_TEMPLATE, out)
    assert out.exists() and out.read_bytes()[:4] == b"%PDF"


def test_docx_introspection():
    pytest.importorskip("docxtpl")
    from docx import Document  # provided by python-docx (docxtpl dep)

    # build a tiny docx template referencing one known + one unknown field
    import tempfile

    doc = Document()
    doc.add_paragraph("{{ engagement_name }} {{ bogus_field }}")
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as fh:
        doc.save(fh.name)
        path = fh.name
    with pytest.raises(ValueError, match="bogus_field"):
        validate_template(path)
